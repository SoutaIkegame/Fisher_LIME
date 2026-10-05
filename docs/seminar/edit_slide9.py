"""Edit slide 9 of the user's seminar4 deck: class-side figure and related-work citations.

Usage: python3 edit_slide9.py <seminar4.pptx> <out.pptx>
"""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path
S=str(Path(__file__).resolve().parent)+'/'
p=Presentation(sys.argv[1])
s=p.slides[8]; s8=p.slides[7]
by={x.shape_id:x for x in s.shapes}
# heading: keep "クラスをまとめる" but drop the claim of novelty
r=by[36].text_frame.paragraphs[0].runs
full=''.join(x.text for x in r)
r[0].text=full.replace('この研究','この研究の方向'); [setattr(x,'text','') for x in r[1:]]
# related-work line in the class box, reusing the style of the note run
note=by[38]
rel=copy.deepcopy(note._element); s.shapes._spTree.append(rel)
from pptx.shapes.autoshape import Shape
rel_sh=[x for x in s.shapes][-1]
rel_sh.top=Inches(6.4); rel_sh.height=Inches(0.4)
note.top=Inches(5.82); by[37].top=Inches(5.12); by[36].top=Inches(4.68)
rr=rel_sh.text_frame.paragraphs[0].runs
rr[0].text='関連研究もクラス間の関係を扱う [2, 4]'; [setattr(x,'text','') for x in rr[1:]]
by[35].height=Inches(2.35); by[35].top=Inches(4.52)
# class-side figure: 10 maps -> arrow -> one axis map
s.shapes.add_picture(S+'class_maps_10.png',Inches(0.35),Inches(5.0),Inches(2.3))
arrow=[x for x in s.shapes if x.shape_id==216][0]
a2=copy.deepcopy(arrow._element); s.shapes._spTree.append(a2)
a2s=[x for x in s.shapes][-1]; a2s.top=Inches(5.3)
s.shapes.add_picture(S+'class_axis_0to6.png',Inches(3.7),Inches(4.62),Inches(1.95),Inches(1.95))
tb=s.shapes.add_textbox(Inches(3.45),Inches(6.6),Inches(2.45),Inches(0.35))
para=tb.text_frame.paragraphs[0]; para.alignment=2
run=para.add_run(); run.text='軸「0から6への移動」'; run.font.size=Pt(14); run.font.color.theme_color=MSO_THEME_COLOR.TEXT_1
tb2=s.shapes.add_textbox(Inches(0.35),Inches(6.25),Inches(2.3),Inches(0.35))
para=tb2.text_frame.paragraphs[0]; para.alignment=2
run=para.add_run(); run.text='クラスごとに10枚'; run.font.size=Pt(14); run.font.color.theme_color=MSO_THEME_COLOR.TEXT_1
# footnote, copied from slide 8's citation box
foot=[x for x in s8.shapes if x.shape_id==266][0]
f2=copy.deepcopy(foot._element); s.shapes._spTree.append(f2)
fs=[x for x in s.shapes][-1]
tf=fs.text_frame
refs=['[2] K. Sokol, P. Flach, LIMETREE: Consistent and Faithful Surrogate Explanations of Multiple Classes, Electronics 14, 929, 2025.（全クラスを1本の多出力木でまとめて説明）',
      '[4] L. Franceschi et al., Explaining Probabilistic Models with Distributional Values, ICML 2024.（特徴ごとにクラスsからrへの遷移を説明）']
paras=tf.paragraphs
for i,pp in enumerate(paras):
    runs=pp.runs
    if i<2:
        runs[0].text=refs[i]; [setattr(x,'text','') for x in runs[1:]]
    else:
        pp._p.getparent().remove(pp._p)
fs.top=Inches(7.06)
p.save(sys.argv[2])  # then replace 、。 with ，． as in edit_user_deck.py
