"""Apply the 10/05 requested edits to the user's hand-edited deck (seminar_3).

Usage: python3 edit_user_deck.py <seminar_3.pptx> <out.pptx> <ten_heatmaps.png>
The uploaded hand-edited deck is not in the repository; the edited result is
docs/seminar/2026-10-05_seminar.pptx.
"""

import re
import sys
import zipfile
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SRC, OUT, HEATMAPS = sys.argv[1:4]
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A}
IN = 914400

prs = Presentation(SRC)
slides = list(prs.slides)
S = lambda n: slides[n - 1]  # 1-based, numbering of the uploaded file


def bbox(sh):
    return sh.left, sh.top, sh.left + sh.width, sh.top + sh.height


def inside(inner, outer, tol=Emu(int(0.05 * IN))):
    a, b = bbox(inner), bbox(outer)
    return a[0] >= b[0] - tol and a[1] >= b[1] - tol and a[2] <= b[2] + tol and a[3] <= b[3] + tol


def is_white(color_parent):
    sc = color_parent.find("a:schemeClr", NS)
    rgb = color_parent.find("a:srgbClr", NS)
    return (sc is not None and sc.get("val") in ("bg1", "lt1")) or (rgb is not None and rgb.get("val").upper() == "FFFFFF")


def is_dark(color_parent):
    sc = color_parent.find("a:schemeClr", NS)
    rgb = color_parent.find("a:srgbClr", NS)
    if sc is not None:
        return sc.get("val") in ("tx1", "dk1", "tx2", "dk2")
    if rgb is not None:
        v = rgb.get("val")
        r, g, b = (int(v[i:i + 2], 16) for i in (0, 2, 4))
        return 0.299 * r + 0.587 * g + 0.114 * b < 100
    return False


def scheme(val):
    el = etree.SubElement(etree.Element(f"{{{A}}}solidFill"), f"{{{A}}}schemeClr")
    el.set("val", val)
    return el.getparent()


# ---------- 1. Outline every remaining filled frame ----------
def outline_frames(slide):
    shapes = list(slide.shapes)
    for sh in shapes:
        sp = sh._element
        spPr = sp.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
        if spPr is None:
            continue
        geom = spPr.find("a:prstGeom", NS)
        if geom is None or geom.get("prst") != "roundRect":
            continue
        if sh.width <= Emu(IN) or sh.height <= Emu(int(0.4 * IN)):
            continue
        fill = spPr.find("a:solidFill", NS)
        if fill is None:
            continue
        dark = is_dark(fill)
        rgb = fill.find("a:srgbClr", NS)
        highlight = rgb is not None and rgb.get("val").upper() == "FBE6DF"
        idx = list(spPr).index(fill)
        spPr.remove(fill)
        spPr.insert(idx, etree.Element(f"{{{A}}}noFill"))
        ln = spPr.find("a:ln", NS)
        if ln is None:
            ln = etree.SubElement(spPr, f"{{{A}}}ln")
        for child in list(ln):
            if child.tag in (f"{{{A}}}solidFill", f"{{{A}}}noFill"):
                ln.remove(child)
        ln.set("w", "28575")
        ln.insert(0, scheme("accent2" if highlight else "tx1"))
        if dark:  # white text that sat on the dark box would vanish
            for other in shapes:
                if other is sh or not other.has_text_frame or not inside(other, sh):
                    continue
                for rPr in other._element.iter(f"{{{A}}}rPr"):
                    f = rPr.find("a:solidFill", NS)
                    if f is not None and is_white(f):
                        rPr.remove(f)
                        rPr.insert(0, scheme("tx1"))


for slide in slides:
    outline_frames(slide)


# ---------- text helpers ----------
def replace_runs(slide, old, new):
    hit = False
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    hit = True
    return hit


def replace_paragraph(slide, old, new):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            full = "".join(r.text for r in p.runs)
            if old in full and p.runs:
                p.runs[0].text = full.replace(old, new)
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                return True
    return False


def add_text(slide, x, y, w, h, runs, size=14, bold=False, color=MSO_THEME_COLOR.TEXT_1, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs if isinstance(runs, list) else [runs]]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for item in para:
            text, opts = (item, {}) if isinstance(item, str) else item
            r = p.add_run()
            r.text = text
            r.font.size = Pt(opts.get("size", size))
            r.font.bold = opts.get("bold", bold)
            r.font.color.theme_color = opts.get("color", color)
    return box


def frame(slide, x, y, w, h, color=MSO_THEME_COLOR.TEXT_1):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.03
    shp.fill.background()
    shp.line.color.theme_color = color
    shp.line.width = Pt(2.25)
    shp.shadow.inherit = False
    return shp


def badge(slide, x, y, label):
    d = 0.45
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    shp.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    shp.shadow.inherit = False
    add_text(slide, x, y, d, d, [(str(label), {"bold": True, "color": MSO_THEME_COLOR.BACKGROUND_1})], size=14, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, col_w, rows, size=12, row_h=0.38):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = gf.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, item in enumerate(row):
            text, opts = (item, {}) if isinstance(item, str) else item
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.theme_color = MSO_THEME_COLOR.TEXT_1 if i == 0 else (MSO_THEME_COLOR.BACKGROUND_1 if i % 2 else MSO_THEME_COLOR.BACKGROUND_2)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = opts.get("align", PP_ALIGN.LEFT)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = i == 0 or opts.get("bold", False)
            r.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1 if i == 0 else opts.get("color", MSO_THEME_COLOR.TEXT_1)
    return gf


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------- 2. Agenda descriptions (deleted slides no longer listed) ----------
for n in (2, 4, 12, 16, 23):
    replace_paragraph(S(n), "手順、上位2クラスとの違い、成り立つための仮定", "手順、成り立つための仮定")
    replace_paragraph(S(n), "仮定の確かめ方、実験の設定、忠実さの測り方、結果", "仮定の確かめ方、実験の設定、結果")
    replace_paragraph(S(n), "課題、今後の計画、まとめ", "課題、今後の計画")

# ---------- 3. Slide 3: the user's ASCII comma ----------
replace_runs(S(3), "だけど , ", "だけど，")
replace_runs(S(3), "だけど ,", "だけど，")

# ---------- 4. Slide 8: citations for the two problems ----------
s8 = S(8)
def append_marker(shape_pred, para_pred, marker):
    for sh in s8.shapes:
        if not (sh.has_text_frame and shape_pred(sh.text_frame.text)):
            continue
        for p in sh.text_frame.paragraphs:
            full = "".join(r.text for r in p.runs)
            if para_pred(full):
                if full.strip() == "()" and p.runs:
                    p.runs[0].text = marker
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)
                else:
                    r = p.add_run(); r.text = " " + marker
                r = p.runs[-1]
                r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
                return True
    return False


def fill_placeholder(marker):
    for sh in s8.shapes:
        if not (sh.has_text_frame and sh.text_frame.text.startswith("実際には10枚")):
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if "()" in r.text:
                    r.text = r.text.replace("()", marker)
                    r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
                    return True
    return False


if not fill_placeholder("[1, 2]"):
    append_marker(lambda t: t.startswith("実際には10枚"), lambda f: f.startswith("確率は高くても"), "[1, 2]")
append_marker(lambda t: "の理由" in t, lambda f: f.startswith("通常のLIME"), "[2, 3]")
refs = [
    [("[1] M. T. Ribeiro, S. Singh, C. Guestrin, “Why Should I Trust You?”: Explaining the Predictions of Any Classifier, KDD 2016.（上位3クラスをクラスごとに別々に説明，Fig. 4）", {})],
    [("[2] K. Sokol, P. Flach, LIMETREE: Consistent and Faithful Surrogate Explanations of Multiple Classes, Electronics 14, 929, 2025.（各クラスの説明を個別に解釈するしかない，Sec. 3）", {})],
    [("[3] P. Nanavati, R. Prasad, CLIMAX: An Exploration of Classifier-Based Contrastive Explanations, IEEE CogMI 2023.（既存手法は予測クラスを正当化するが，他のクラスとの違いは保証しない）", {})],
]
add_text(s8, 0.6, 6.74, 11.4, 0.66, refs, size=8, color=MSO_THEME_COLOR.TEXT_1)

# ---------- 5. Slide 11: ten different heatmaps ----------
s11 = S(11)
for sh in list(s11.shapes):
    x, y = sh.left / IN, sh.top / IN
    if 0.85 <= x <= 3.4 and 2.35 <= y <= 5.6 and sh.width < Emu(int(0.12 * IN)):
        sh._element.getparent().remove(sh._element)
s11.shapes.add_picture(HEATMAPS, Inches(0.78), Inches(2.4), Inches(3.3), Inches(3.3 * 3.0 / 3.1))

# ---------- 6. Slide 22: coefficient count ----------
replace_paragraph(S(22), "係数は640個から222個に減る（手書き数字）", "係数は640個から192個に減る（手書き数字）")

# Slide 22 (結果2): the user's bullet box overflowed by one line; make it and its frame taller.
for sh in S(22).shapes:
    if sh.has_text_frame and "SVM" in sh.text_frame.text and sh.left > Emu(int(7 * IN)):
        box = sh
        for other in S(22).shapes:
            if other is not box and abs(other.left - box.left) < Emu(int(0.3 * IN)) and abs(other.top - box.top) < Emu(int(0.3 * IN)):
                other.height = other.height + Emu(int(0.35 * IN))
        box.height = box.height + Emu(int(0.35 * IN))
        break

# ---------- 7. Slide 24: rebuild 課題 ----------
s24 = S(24)
for sh in list(s24.shapes):
    if sh.is_placeholder or (sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > Emu(int(6.9 * IN))):
        continue
    sh._element.getparent().remove(sh._element)

frame(s24, 0.6, 1.3, 6.15, 5.45)
badge(s24, 0.82, 1.47, 1)
add_text(s24, 1.42, 1.47, 5.2, 0.45, "軸が読めるとは限らない", size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
C = PP_ALIGN.CENTER
table(s24, 0.85, 2.15, 5.65, [1.55, 2.85, 1.25], [
    ["軸の形", "例（重みの大きいクラス）", ("読めるか", {"align": C})],
    ["2クラスの対比", "6 + ／ 0 −", ("読める", {"align": C, "bold": True, "color": MSO_THEME_COLOR.ACCENT_1})],
    ["1対多", "6 +0.8 ／ 0 −0.5, 9 −0.3", ("読める", {"align": C, "bold": True, "color": MSO_THEME_COLOR.ACCENT_1})],
    ["混在", "6 +0.37, 1 +0.21 ／ 0 −0.46", ("読めない", {"align": C, "bold": True, "color": MSO_THEME_COLOR.ACCENT_2})],
], size=12, row_h=0.45)
add_text(s24, 0.85, 4.1, 5.7, 1.0, [
    [("1対多なら「6の確率は主に0から，一部は9から来た」と読める．混在だと，6と1が同じ側にあり意味が取れない", {})],
], size=13)
add_text(s24, 0.85, 5.05, 5.7, 0.5, [[("実際の軸の例（手書き数字）　2 −0.76，3 +0.59，9 +0.27（1対多の形）", {})]], size=12, color=MSO_THEME_COLOR.TEXT_1)
add_text(s24, 0.85, 5.75, 5.7, 0.8, "次元が減っても，軸の意味が分からなければ解釈しやすくならない", size=13, bold=True, color=MSO_THEME_COLOR.ACCENT_2)

frame(s24, 6.95, 1.3, 5.75, 2.55)
badge(s24, 7.15, 1.47, 2)
add_text(s24, 7.75, 1.47, 4.85, 0.45, "多くのクラスがまとまって動く場合は未検証", size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE)
add_text(s24, 7.2, 2.07, 5.3, 1.75, [
    [("今回のデータでは動くクラスが3〜4個で，まとまって動いても軸が1本減る程度（yeast 45%，手書き数字 32%の近傍）", {})],
    [("検証案　似たクラスが群になったデータ（100クラスが20の上位クラスにまとまるCIFAR-100など）で，動くクラスの数と必要な軸の数を比べる", {})],
], size=12)

frame(s24, 6.95, 4.05, 5.75, 2.7)
badge(s24, 7.15, 4.22, 3)
add_text(s24, 7.75, 4.22, 4.8, 0.45, "ランダムフォレストではうまくいかない", size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
table(s24, 7.2, 4.85, 5.3, [1.65, 1.05, 1.0, 1.6], [
    ["", ("動くクラス（個）", {"align": C}), ("必要な軸（本）", {"align": C}), ("3本の R²（通常 → この手法）", {"align": C})],
    ["MLP", ("3.15", {"align": C}), ("2.09", {"align": C}), ("0.793 → 0.785", {"align": C})],
    ["ランダムフォレスト", ("14.31", {"align": C}), ("11.06", {"align": C, "bold": True, "color": MSO_THEME_COLOR.ACCENT_2}), ("0.638 → 0.519", {"align": C, "bold": True, "color": MSO_THEME_COLOR.ACCENT_2})],
], size=10, row_h=0.42)
add_text(s24, 7.2, 6.18, 5.35, 0.55, [[("合成データ，20クラス，半径0.15．出力が木の多数決で階段状（ガタガタ）なので，多くのクラスが少しずつ揺れ，少ない軸にも直線にも収まらない", {})]], size=10)

set_notes(s24, "課題は3つあります．1つ目は軸の読みやすさです．PCAの軸は，複数のクラスに重みが分かれることがあります．6が上がって0と9が下がる形なら「6の確率は主に0から来た」と読めますが，6と1が同じ側に並ぶと意味が取れません．手書き数字の実際の軸は「2が下がり，3と9が上がる」という1対多の形でした．2つ目は，多くのクラスがまとまって動く場合をまだ検証していないことです．今回のデータでは動くクラスが3〜4個しかなく，まとまって動いても軸が1本減る程度でした．似たクラスが群になったCIFAR-100のようなデータで確かめる予定です．3つ目は，ランダムフォレストではうまくいかないことです．合成データの20クラスでは，動くクラスが14個，必要な軸が11本で，3本にまとめるとR²が0.64から0.52に落ちました．ランダムフォレストの出力は木の多数決で階段状になるので，多くのクラスが少しずつ揺れ，少ない軸にも直線にも収まりません．")

# ---------- 8. Delete slides 14, 20, 26 and the appendix (27-33) ----------
remove = {14, 20, 26} | set(range(27, 34))
sldIdLst = prs.slides._sldIdLst
for i, sldId in reversed(list(enumerate(list(sldIdLst), start=1))):
    if i in remove:
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)

prs.save(OUT)

# ---------- 9. Comma and period instead of 、 and 。 (slides, notes, charts) ----------
with zipfile.ZipFile(OUT) as zin:
    items = [(i, zin.read(i.filename)) for i in zin.infolist()]
with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
    for info, data in items:
        name = info.filename
        if re.match(r"ppt/(slides|notesSlides|charts)/[^/]+\.xml$", name):
            text = data.decode("utf-8").replace("、", "，").replace("。", "．")
            data = text.encode("utf-8")
        zout.writestr(info, data)
print("saved", OUT)
